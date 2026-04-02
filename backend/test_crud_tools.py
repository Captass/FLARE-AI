import asyncio
import os
import sys

# Ajouter le chemin parent pour que 'backend' soit reconnu comme package
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '..')))

from langchain_core.runnables import RunnableConfig

async def test_crud_word():
    print("Testing Word CRUD tools...")
    from backend.agents.workers.document_worker import (
        generate_word_document, 
        read_word_document_as_json, 
        read_word_document, 
        update_word_document, 
        delete_word_document,
        DocumentWorker
    )
    worker = DocumentWorker()
    print("DocumentWorker initialized tools:", [t.name for t in worker.tools])
    
    # 1. Create
    filename = "test_crud_word.docx"
    elements_json = '[{"type": "paragraph", "text": "Test CREATE."}]'
    config = RunnableConfig(configurable={"user_id": "test_user", "session_id": "test_session"})
    
    create_result = await generate_word_document.ainvoke({"filename": filename, "elements_json": elements_json}, config=config)
    print("CREATE RESULT:", create_result)
    
    # Normally create_result contains URL or path. Since we might not have firebase configured properly in test,
    # let's assume local test file to read, update, delete.
    
    # Since generate_word_document handles saving locally to memory and uploading to firebase.
    # It might fail uploading if not configured but will still succeed locally in doc_bytes.
    # Let's mock a file_path_or_url
    
    with open("test_local.docx", "wb") as f:
        from docx import Document
        doc = Document()
        doc.add_paragraph("Local test")
        doc.save(f)
        
    read_res = await read_word_document.ainvoke({"file_path_or_url": "test_local.docx"}, config=config)
    print("READ RESULT:", read_res)
    
    read_json_res = await read_word_document_as_json.ainvoke({"file_path_or_url": "test_local.docx"}, config=config)
    print("READ JSON RESULT:", read_json_res)
    
    update_res = await update_word_document.ainvoke({
        "file_path_or_url": "test_local.docx", 
        "elements_json": '[{"type": "paragraph", "text": "Updated."}]'
    }, config=config)
    print("UPDATE RESULT:", update_res)
    
    delete_res = await delete_word_document.ainvoke({"file_path_or_url": "test_local.docx"}, config=config)
    print("DELETE RESULT:", delete_res)


async def test_crud_excel():
    print("\nTesting Excel CRUD tools...")
    from backend.agents.workers.spreadsheet_worker import (
        generate_excel_document, 
        read_excel_document_as_json, 
        read_excel_document, 
        update_excel_document, 
        delete_excel_document,
        SpreadsheetWorker
    )
    worker = SpreadsheetWorker()
    print("SpreadsheetWorker initialized tools:", [t.name for t in worker.tools])
    
    config = RunnableConfig(configurable={"user_id": "test_user", "session_id": "test_session"})
    
    with open("test_local.xlsx", "wb") as f:
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "Local test"
        wb.save(f)
        
    read_res = await read_excel_document.ainvoke({"file_path_or_url": "test_local.xlsx"}, config=config)
    print("READ RESULT:", read_res)
    
    read_json_res = await read_excel_document_as_json.ainvoke({"file_path_or_url": "test_local.xlsx"}, config=config)
    print("READ JSON RESULT:", read_json_res)
    
    update_res = await update_excel_document.ainvoke({
        "file_path_or_url": "test_local.xlsx", 
        "sheets_json": '[{"name": "Sheet1", "rows": [[{"value": "Updated"}]]}]'
    }, config=config)
    print("UPDATE RESULT:", update_res)
    
    delete_res = await delete_excel_document.ainvoke({"file_path_or_url": "test_local.xlsx"}, config=config)
    print("DELETE RESULT:", delete_res)

async def main():
    await test_crud_word()
    await test_crud_excel()

if __name__ == "__main__":
    asyncio.run(main())
